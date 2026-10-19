import zipfile
import xml.etree.ElementTree as ET
import openpyxl

filepath = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=True)
ws = wb["LURICANTE AROIL"]

with zipfile.ZipFile(filepath, "r") as z:
    draw_root = ET.fromstring(z.read("xl/drawings/drawing9.xml"))
    draw_rels = ET.fromstring(z.read("xl/drawings/_rels/drawing9.xml.rels"))
    rel_map = {r.attrib["Id"]: r.attrib["Target"] for r in draw_rels}
    
    ns = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
    }
    
    # Map images to (center_row, col)
    assigned = {}
    for anchor in draw_root:
        from_el = anchor.find("xdr:from", ns)
        to_el = anchor.find("xdr:to", ns)
        if from_el is None:
            continue
        c1 = int(from_el.find("xdr:col", ns).text) + 1
        r1 = int(from_el.find("xdr:row", ns).text) + 1
        c2 = int(to_el.find("xdr:col", ns).text) + 1 if to_el is not None else c1
        r2 = int(to_el.find("xdr:row", ns).text) + 1 if to_el is not None else r1
        
        center_row = round((r1 + r2) / 2.0)
        center_col = c1
        
        pic = anchor.find("xdr:pic", ns)
        if pic is None:
            continue
        blip = pic.find(".//a:blip", ns)
        if blip is None:
            continue
        embed = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        
        if center_row not in assigned:
            assigned[center_row] = {}
        assigned[center_row][center_col] = (embed, rel_map.get(embed))
        
    print("Aroil Row Mapping by Center Row:")
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        cols_data = assigned.get(r, {})
        print(f"Row {r}: '{name}' -> Cols: {list(cols_data.keys())}")
