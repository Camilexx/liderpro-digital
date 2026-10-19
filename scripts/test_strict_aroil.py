import openpyxl, os, re

folder = r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES'
wb = openpyxl.load_workbook(os.path.join(folder, 'FILTROS.xlsx'), data_only=True)
ws = wb['LURICANTE AROIL']

inv_wb = openpyxl.load_workbook(os.path.join(folder, 'INVENTARIO.xlsx'), data_only=True)
inv_ws = inv_wb.active
inv_items = []
for r in range(8, inv_ws.max_row + 1):
    sku = str(inv_ws.cell(r, 1).value or '').strip()
    name = str(inv_ws.cell(r, 2).value or '').strip()
    stock = inv_ws.cell(r, 6).value
    if sku and name:
        inv_items.append({'sku': sku, 'name': name, 'stock': stock})

for r in range(2, ws.max_row + 1):
    c2 = str(ws.cell(r, 2).value or '').strip()
    if not c2: continue
    print(f"\nRow {r:2d}: {c2}")
    
    matches = []
    c2_clean = c2.upper()
    for it in inv_items:
        it_name = it['name'].upper()
        if 'AROIL' not in it_name: continue
        
        # Exact product line matching
        if 'RACING 10W30' in c2_clean:
            if 'RACING 10W30' in it_name: matches.append(it)
        elif 'RACING 20W50' in c2_clean:
            if 'RACING 20W50' in it_name: matches.append(it)
        elif 'SUPER 40' in c2_clean:
            if 'SUPER 40' in it_name: matches.append(it)
        elif '15W40 CI-4' in c2_clean:
            if '15W40 CI-4' in it_name: matches.append(it)
        elif '25W60 CF' in c2_clean:
            if '25W60 CF' in it_name: matches.append(it)
        elif 'DIESEL 40 CF' in c2_clean:
            if 'DIESEL 40 CF' in it_name: matches.append(it)
        elif 'MOTOR 2T' in c2_clean:
            if '2T' in it_name: matches.append(it)
        elif '4T 20W50' in c2_clean:
            if '4T 20W50' in it_name: matches.append(it)
        elif '75W85 GL-5' in c2_clean:
            if '75W85' in it_name: matches.append(it)
        elif '80W90 GL-5' in c2_clean:
            if '80W90' in it_name: matches.append(it)
        elif '85W140 GL-5' in c2_clean:
            if '85W140' in it_name: matches.append(it)
        elif '90 GL-4' in c2_clean:
            # Must have 90 GL-4 and NOT 80W90 or 140
            if ('90 GL-4' in it_name or '90 GL' in it_name) and '80W' not in it_name and '140' not in it_name and '75W' not in it_name:
                matches.append(it)
        elif '140 GL-4' in c2_clean:
            # Must have 140 and NOT 85W140
            if ('140 GL-4' in it_name or '140 GL' in it_name) and '85W' not in it_name:
                matches.append(it)
        elif 'ATF DIII' in c2_clean:
            if 'ATF' in it_name: matches.append(it)
        elif 'AW ISO 68 ROJO' in c2_clean:
            if 'ISO 68 ROJO' in it_name or 'AW ISO 68 ROJO' in it_name: matches.append(it)
        elif 'TRACFLUID 10W30' in c2_clean:
            if 'TRACFLUID' in it_name: matches.append(it)
        elif 'ISO 68 AMBAR' in c2_clean:
            if 'ISO 68 AMBAR' in it_name: matches.append(it)
        elif 'GRASA CHASSIS 2 ROJA' in c2_clean:
            if 'GRASA' in it_name and 'ROJA' in it_name: matches.append(it)
        elif 'GRASA CHASSIS 2 AZUL' in c2_clean:
            if 'GRASA' in it_name and 'AZUL' in it_name: matches.append(it)
            
    for m in matches:
        print(f"   -> SKU: {m['sku']} | Stk: {m['stock']} | {m['name']}")
    if not matches:
        print("   -> NO MATCH FOUND")
