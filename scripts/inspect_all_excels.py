import openpyxl
import os

files = [
    r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS DE AIRE AA.xlsx',
    r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx',
    r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FOTOS MARCAS.xlsx',
    r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\INVENTARIO ACEITES.xlsx',
    r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\INVENTARIO.xlsx'
]

for fpath in files:
    fname = os.path.basename(fpath)
    wb = openpyxl.load_workbook(fpath, read_only=True)
    print(f"\n==========================================")
    print(f"FILE: {fname}")
    print(f"SHEETS: {wb.sheetnames}")
    for sname in wb.sheetnames:
        sheet = wb[sname]
        print(f"   Sheet '{sname}': max_row={sheet.max_row}, max_col={sheet.max_column}")
