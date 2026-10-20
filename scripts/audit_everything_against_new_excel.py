import openpyxl, json, os

excel_path = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\INVENTARIO FOTOS.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

base_dir = r"c:\Users\PC\Desktop\NEXUS\liderpro-digital"
with open(os.path.join(base_dir, "src/data/products.ts"), "r", encoding="utf-8") as f:
    content = f.read()

start = content.find("export const PRODUCTS: Product[] = ")
json_str = content[start + len("export const PRODUCTS: Product[] = "):].rstrip(";\n ")
products = json.loads(json_str)

prod_by_sku = {p["sku"]: p for p in products}
prod_names_lower = {p["name"].lower().strip(): p for p in products}

print(f"Total products in src/data/products.ts: {len(products)}")

# Check each sheet in Excel
sheets = [s for s in wb.sheetnames if s != "Sheet1"]
total_excel_items = 0

print("\n=== VERIFYING EACH EXCEL SHEET ITEMS ===")
for sheetname in sheets:
    ws = wb[sheetname]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(x is not None and str(x).strip() != '' for x in r)]
    total_excel_items += len(non_empty)
    
    # Check match rate
    matched = 0
    missing = []
    for r in non_empty:
        # Find product name and code
        p_name = None
        code = None
        for c in r:
            if c and isinstance(c, str):
                c_clean = c.strip()
                if len(c_clean) > 2:
                    if not p_name:
                        p_name = c_clean
                    elif not code and any(char.isdigit() for char in c_clean):
                        code = c_clean
        
        # Match by name or partial
        found = False
        if p_name:
            p_clean = p_name.lower().replace('"', '').replace('\t', ' ').strip()
            for prod in products:
                prod_clean = prod["name"].lower().replace('"', '').replace('\t', ' ').strip()
                if p_clean in prod_clean or prod_clean in p_clean:
                    found = True
                    break
        if found:
            matched += 1
        else:
            missing.append(p_name)
            
    print(f"Sheet: {sheetname:22} | Excel items: {len(non_empty):2} | Matched in catalog: {matched:2} | Unmatched: {len(missing)}")
    if missing and len(missing) <= 5:
        for m in missing:
            print(f"   Missing item: {m}")

print(f"\nTotal items across sheets: {total_excel_items}")
