import openpyxl, os, zipfile
from xml.etree import ElementTree as ET

excel_path = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\INVENTARIO FOTOS.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("=== SHEET SUMMARY ===")
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(x is not None and str(x).strip() != '' for x in r)]
    max_c = ws.max_column
    max_r = ws.max_row
    print(f"Sheet: {sheetname:25} | Max Col: {max_c:2} | Max Row: {max_r:4} | Non-empty rows: {len(non_empty)}")

print("\n=== DETAIL ON BATERIAS BOSCH ===")
ws_bb = wb['BATERIAS BOSCH']
for idx, r in enumerate(ws_bb.iter_rows(values_only=True), 1):
    if any(x is not None for x in r):
        print(f"Row {idx:2}: {[x for x in r if x is not None]}")

print("\n=== DETAIL ON BATERIAS ECUADOR ===")
ws_be = wb['BATERIAS ECUADOR']
for idx, r in enumerate(ws_be.iter_rows(values_only=True), 1):
    if any(x is not None for x in r):
        print(f"Row {idx:2}: {[x for x in r if x is not None]}")
