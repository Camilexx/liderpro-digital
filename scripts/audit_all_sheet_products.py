import zipfile
import xml.etree.ElementTree as ET
import openpyxl
import os

filepath = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=True)

with zipfile.ZipFile(filepath, "r") as z:
    wb_root = ET.fromstring(z.read("xl/workbook.xml"))
    wb_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    sheet_target = {r.attrib["Id"]: r.attrib["Target"] for r in wb_rels}
    
    ns_wb = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    ns_draw = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
    }
    
    total_products_with_images = 0
    total_products_without_images = 0
    
    for sheetname in wb.sheetnames:
        if sheetname == "Sheet1":
            continue
        ws = wb[sheetname]
        
        target = None
        for s in wb_root.findall(".//s:sheet", ns_wb):
            if s.attrib["name"] == sheetname:
                target = sheet_target[s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
                break
                
        rel_file = f"xl/worksheets/_rels/{target.split('/')[-1]}.rels"
        if rel_file not in z.namelist():
            print(f"Sheet {sheetname}: no rel file")
            continue
            
        s_rels = ET.fromstring(z.read(rel_file))
        draw_targets = [r.attrib["Target"] for r in s_rels if "drawing" in r.attrib["Type"]]
        if not draw_targets:
            print(f"Sheet {sheetname}: no drawing target")
            continue
            
        draw_target = draw_targets[0]
        draw_filename = "xl/drawings/" + draw_target.split("/")[-1]
        draw_rels_file = "xl/drawings/_rels/" + draw_target.split("/")[-1] + ".rels"
        
        dr_rels_root = ET.fromstring(z.read(draw_rels_file))
        rel_map = {r.attrib["Id"]: r.attrib["Target"] for r in dr_rels_root}
        draw_root = ET.fromstring(z.read(draw_filename))
        
        anchors_by_row = {}
        for anchor in draw_root:
            from_el = anchor.find("xdr:from", ns_draw)
            if from_el is None:
                continue
            col = int(from_el.find("xdr:col", ns_draw).text) + 1
            row = int(from_el.find("xdr:row", ns_draw).text) + 1
            blip = anchor.find(".//a:blip", ns_draw)
            if blip is not None:
                embed = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if row not in anchors_by_row:
                    anchors_by_row[row] = {}
                anchors_by_row[row][col] = rel_map.get(embed)
                
        sheet_with_img = 0
        sheet_no_img = 0
        
        for r in range(1, ws.max_row + 1):
            val_a = ws.cell(r, 1).value
            val_b = ws.cell(r, 2).value
            
            # If both are None or empty, skip
            if not val_a and not val_b:
                continue
            if str(val_b or val_a).strip() == "":
                continue
                
            imgs = anchors_by_row.get(r, {})
            if imgs:
                sheet_with_img += 1
            else:
                sheet_no_img += 1
                
        total_products_with_images += sheet_with_img
        total_products_without_images += sheet_no_img
        print(f"Sheet '{sheetname}': {sheet_with_img} with images, {sheet_no_img} without images")

    print(f"\nTOTAL: {total_products_with_images} products with images, {total_products_without_images} without images")
