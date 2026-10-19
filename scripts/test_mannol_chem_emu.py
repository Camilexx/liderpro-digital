import zipfile
import xml.etree.ElementTree as ET
import openpyxl

filepath = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=True)
ws = wb["MANNOL"]

cumulative_y = [0]
default_height_pt = ws.sheet_format.defaultRowHeight or 15.0
total_y = 0
for r in range(1, ws.max_row + 10):
    h = ws.row_dimensions[r].height if r in ws.row_dimensions and ws.row_dimensions[r].height else default_height_pt
    h_emu = int(round(h * 12700))
    total_y += h_emu
    cumulative_y.append(total_y)

with zipfile.ZipFile(filepath, "r") as z:
    draw_root = ET.fromstring(z.read("xl/drawings/drawing10.xml"))
    draw_rels = ET.fromstring(z.read("xl/drawings/_rels/drawing10.xml.rels"))
    rel_map = {r.attrib["Id"]: r.attrib["Target"] for r in draw_rels}
    ns = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
    }
    
    anchors_info = []
    for anchor in draw_root:
        from_el = anchor.find("xdr:from", ns)
        to_el = anchor.find("xdr:to", ns)
        if from_el is None:
            continue
        c1 = int(from_el.find("xdr:col", ns).text) + 1
        r1 = int(from_el.find("xdr:row", ns).text) + 1
        off_y1 = int(from_el.find("xdr:rowOff", ns).text if from_el.find("xdr:rowOff", ns) is not None else 0)
        c2 = int(to_el.find("xdr:col", ns).text) + 1 if to_el is not None else c1
        r2 = int(to_el.find("xdr:row", ns).text) + 1 if to_el is not None else r1
        off_y2 = int(to_el.find("xdr:rowOff", ns).text if to_el is not None and to_el.find("xdr:rowOff", ns) is not None else 0)
        
        pic = anchor.find("xdr:pic", ns)
        if pic is None:
            continue
        blip = pic.find(".//a:blip", ns)
        if blip is None:
            continue
        embed = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        
        y_start = cumulative_y[r1 - 1] + off_y1
        y_end = cumulative_y[r2 - 1] + off_y2
        y_mid = (y_start + y_end) / 2.0
        
        assigned_row = 1
        for r in range(1, len(cumulative_y)):
            if cumulative_y[r-1] <= y_mid < cumulative_y[r]:
                assigned_row = r
                break
        anchors_info.append((assigned_row, c1, embed, rel_map.get(embed), y_mid))

by_row = {}
for r, c, emb, tgt, y_mid in anchors_info:
    if r not in by_row:
        by_row[r] = {}
    by_row[r][c] = tgt

print("=== EXACT EMU-BASED MAPPING FOR MANNOL (CHEMICALS/CARE) ===")
for r in range(2, min(ws.max_row + 1, 15)):
    name = ws.cell(r, 2).value or ws.cell(r, 1).value
    if name:
        print(f"Row {r}: '{name.strip()}' -> Cols: {sorted(by_row.get(r, {}).keys())} -> {by_row.get(r, {})}")
