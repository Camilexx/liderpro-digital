import openpyxl
import os
import re

folder = r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES'
wb = openpyxl.load_workbook(os.path.join(folder, 'FILTROS.xlsx'), data_only=True)

sheets = [
    'LURICANTE MOBIL', 'LURICANTE AMALIE', 'LURICANTE MOTUL', 
    'LURICANTE VALVOLINE', 'LURICANTE AROIL', 'LURICANTE MANNOL', 
    'ABRO', 'PERFUMES', 'MANNOL'
]

for s in sheets:
    ws = wb[s]
    print(f"\n=======================================================")
    print(f"SHEET: {s} (max_row={ws.max_row})")
    print(f"=======================================================")
    for r in range(1, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        c2 = ws.cell(r, 2).value
        if c1 or c2:
            print(f"Row {r:2d}: C1={repr(c1):<20} | C2={repr(c2)}")
