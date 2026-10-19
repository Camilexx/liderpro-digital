import openpyxl
import os

inv_path = r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\INVENTARIO.xlsx'
wb = openpyxl.load_workbook(inv_path, data_only=True)
sheet = wb.active
print(f"Sheet title: {sheet.title}, Max rows: {sheet.max_row}, Max cols: {sheet.max_column}")

for r in range(1, min(15, sheet.max_row + 1)):
    row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    print(f"Row {r}: {row_vals[:6]}")
