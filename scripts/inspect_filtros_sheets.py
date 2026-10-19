import openpyxl
import os

filepath = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=True)

print("Sheets in FILTROS.xlsx:")
for sname in wb.sheetnames:
    ws = wb[sname]
    max_r = ws.max_row
    max_c = ws.max_column
    
    # Check non-empty rows
    sample_rows = []
    for r in range(1, min(max_r + 1, 30)):
        vals = [ws.cell(r, c).value for c in range(1, min(max_c + 1, 10))]
        if any(v is not None and str(v).strip() != "" for v in vals):
            sample_rows.append((r, vals[:6]))
    
    print(f"\n--- Sheet: '{sname}' (max_row={max_r}, max_col={max_c}, sample rows={len(sample_rows)}) ---")
    for r, vals in sample_rows[:8]:
        print(f"  Row {r}: {vals}")
