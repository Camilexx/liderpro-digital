import openpyxl, os, re

folder = r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES'
wb = openpyxl.load_workbook(os.path.join(folder, 'FILTROS.xlsx'), data_only=True)
ws = wb['LURICANTE MANNOL']

inv_wb = openpyxl.load_workbook(os.path.join(folder, 'INVENTARIO.xlsx'), data_only=True)
inv_ws = inv_wb.active
inv_items = []
for r in range(8, inv_ws.max_row + 1):
    sku = str(inv_ws.cell(r, 1).value or '').strip()
    name = str(inv_ws.cell(r, 2).value or '').strip()
    stock = inv_ws.cell(r, 6).value
    if sku and name:
        inv_items.append({'sku': sku, 'name': name, 'stock': stock})

for r in range(2, ws.max_row + 1):
    c2 = str(ws.cell(r, 2).value or '').strip()
    if not c2: continue
    print(f"\nRow {r:2d}: {c2}")
    
    # Extract viscosity / product code from c2
    # Find matching items in inv_items that belong strictly to this specific product
    matches = []
    for it in inv_items:
        # Check brand
        if 'MANNOL' not in it['name'].upper():
            continue
        
        # Test specific product matches
        # If c2 is "MANNOL REFRIGERANTE AZUL" -> AG 11
        if 'AZUL' in c2.upper() and 'REFRIGERANTE' in c2.upper():
            if 'AZUL' in it['name'].upper() and 'REFRIGERANTE' in it['name'].upper():
                matches.append(it)
        elif 'ROJO' in c2.upper() and 'REFRIGERANTE' in c2.upper():
            if 'ROJO' in it['name'].upper() and 'REFRIGERANTE' in it['name'].upper():
                matches.append(it)
        elif 'VERDE' in c2.upper() and 'REFRIGERANTE' in c2.upper():
            if 'VERDE' in it['name'].upper() and 'REFRIGERANTE' in it['name'].upper():
                matches.append(it)
        elif 'NARANJA' in c2.upper() and 'REFRIGERANTE' in c2.upper():
            if 'NARANJA' in it['name'].upper() and 'REFRIGERANTE' in it['name'].upper():
                matches.append(it)
        else:
            # Check for specific grade/model
            # E.g. TS-4, SAFARI 20W50, CLASSIC 10W40, DEFENDER, SPECIAL PLUS 10W30, TOYOTA LEXUS, HYUNDAI KIA, CERAMIC 5W30, etc.
            c2_clean = re.sub(r'[\t\r\n]', ' ', c2.upper())
            # Find key identifiers
            for ident in [
                'TS-4', 'SAFARI', 'CLASSIC 10W40', 'DEFENDER', 'SPECIAL PLUS',
                'TOYOTA LEXUS', 'HYUNDAI KIA', 'CERAMIC PRO', 'CERAMIC 5W30',
                '4TAKT MOTORBIKE 10W40', '4TAKT MOTORBIKE 10W50', 'POWERBIKE',
                'DIESEL TURBO 5W40', 'ENERGY ULTRA JP 5W20', 'DIESEL TDI 5W30',
                'ENERGY FORMULA JP 5W30', 'FOR CHINESE CARS', 'FWD 75W85',
                'HYPOID LSD 85W140', 'HYPOID GETRIEBEOEL 80W90', 'UNIVERSAL GETRIEBEOEL 80W90',
                'BASIC PLUS 75W90', 'MTF-3 75W', 'DCT FLUID', 'ATF DEXRON III',
                'ATF DEXRON VI', 'CVT NISSAN NS-3', 'ATF ULV', 'HYDRO HV ISO 32',
                'HYDRO HV ISO 46', 'HYDRO HV ISO 68', 'MULTIFARM STOU',
                'MOLIBDEN GREEN ULTRA 5W30', 'TS-5 UHPD'
            ]:
                if ident in c2_clean:
                    # Match only if ident is in it['name']
                    # Normalize ident words
                    id_words = ident.split()
                    if all(iw in it['name'].upper() for iw in id_words):
                        # Ensure we don't cross 10W40 with 10W50 or ISO 32 with ISO 46/68 or CERAMIC with CERAMIC PRO
                        if 'CERAMIC' in ident and 'PRO' not in ident and 'PRO' in it['name'].upper():
                            continue
                        if '10W40' in ident and '10W50' in it['name'].upper():
                            continue
                        if '10W50' in ident and '10W40' in it['name'].upper():
                            continue
                        if 'ISO 32' in ident and ('ISO 46' in it['name'].upper() or 'ISO 68' in it['name'].upper()):
                            continue
                        if 'ISO 46' in ident and ('ISO 32' in it['name'].upper() or 'ISO 68' in it['name'].upper()):
                            continue
                        if 'ISO 68' in ident and ('ISO 32' in it['name'].upper() or 'ISO 46' in it['name'].upper()):
                            continue
                        matches.append(it)
                    break
    for m in matches:
        print(f"   -> SKU: {m['sku']} | Stk: {m['stock']} | {m['name']}")
    if not matches:
        print("   -> NO MATCH FOUND")
