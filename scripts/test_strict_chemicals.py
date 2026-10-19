import openpyxl, os, re

folder = r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES'
wb = openpyxl.load_workbook(os.path.join(folder, 'FILTROS.xlsx'), data_only=True)

inv_wb = openpyxl.load_workbook(os.path.join(folder, 'INVENTARIO.xlsx'), data_only=True)
inv_ws = inv_wb.active
inv_items = []
for r in range(8, inv_ws.max_row + 1):
    sku = str(inv_ws.cell(r, 1).value or '').strip()
    name = str(inv_ws.cell(r, 2).value or '').strip()
    stock = inv_ws.cell(r, 6).value
    if sku and name:
        inv_items.append({'sku': sku, 'name': name, 'stock': stock})

for sname in ['ABRO', 'MANNOL']:
    ws = wb[sname]
    print(f"\n=================== {sname} ===================")
    for r in range(2, ws.max_row + 1):
        c2 = str(ws.cell(r, 2).value or '').strip()
        if not c2: continue
        
        matches = []
        c2_clean = c2.upper()
        
        for it in inv_items:
            it_name = it['name'].upper()
            
            # Specific match logic
            if sname == 'ABRO':
                if 'ABRO' not in it_name: continue
                if 'AB-404' in c2_clean and 'AB-404' in it_name: matches.append(it)
                elif 'AB-80' in c2_clean and 'AB-80' in it_name: matches.append(it)
                elif 'REFRIGERANTE ROJO' in c2_clean and 'REFRIGERANTE ROJO' in it_name: matches.append(it)
                elif 'REFRIGERANTE VERDE' in c2_clean and 'REFRIGERANTE VERDE' in it_name: matches.append(it)
                elif 'LG-990' in c2_clean and 'LG-990' in it_name: matches.append(it)
                elif 'SS-822' in c2_clean and 'SS-822' in it_name: matches.append(it)
                elif 'TAIPE' in c2_clean and 'TAIPE' in it_name: matches.append(it)
                elif 'TRANSPARENTE' in c2_clean and 'TRANSPARENTE' in it_name: matches.append(it)
            elif sname == 'MANNOL':
                if 'MANNOL' not in it_name: continue
                
                # Check specific model numbers or distinct keywords
                # e.g. 9929 (ESTER ADDITIVE ACEITE), 9930 (DIESEL ESTER), 9950 (BENZIN ESTER),
                # 9966 (RADIADOR), 9968 (CAJA DE CAMBIOS), 9980 (INYECTORES DIESEL), 9981 (INYECTORES GASOLINA)
                if 'TRATAMIENTO DE ACEITE' in c2_clean or 'ESTER ADDITIVE / TRATAMIENTO' in c2_clean:
                    if 'ESTER ADDITIVE / TRATAMIENTO DE ACEITE' in it_name: matches.append(it)
                elif 'DIESEL ESTER ADDITIVE' in c2_clean:
                    if 'DIESEL ESTER ADDITIVE' in it_name: matches.append(it)
                elif 'BENZIN ESTER ADDITIVE' in c2_clean:
                    if 'BENZIN ESTER' in it_name: matches.append(it)
                elif 'SELLADOR DE FUGAS PARA RADIADOR' in c2_clean:
                    if 'RADIADOR' in it_name and 'SELLADOR' in it_name: matches.append(it)
                elif 'SELLADOR DE FUGAS CAJA' in c2_clean:
                    if 'CAJA DE CAMBIOS' in it_name and 'SELLADOR' in it_name: matches.append(it)
                elif 'LIMPIADOR INYECTORES DIESEL' in c2_clean:
                    if 'INYECTORES DIESEL' in it_name and 'LIMPIADOR' in it_name: matches.append(it)
                elif 'LIMPIA INYECTORES GASOLINA' in c2_clean:
                    if 'INYECTORES GASOLINA' in it_name and 'LIMPIA' in it_name: matches.append(it)
                elif 'CAJA AUTOM' in c2_clean:
                    if 'AUTOM' in it_name and 'CAJA' in it_name: matches.append(it)
                elif 'CAJA MANUAL' in c2_clean:
                    if 'MANUAL' in it_name and 'CAJA' in it_name: matches.append(it)
                else:
                    # Match by unique title substring
                    title_words = [w for w in re.split(r'[\s\-\/\(\)\,]+', c2_clean) if len(w) >= 3 and w not in ['MANNOL', 'PARA', 'CON', 'DEL', 'LTS', 'LT', 'UND', 'CJ', 'X']]
                    if title_words and all(tw in it_name for tw in title_words[:3]):
                        matches.append(it)

        if len(matches) > 0:
            print(f"Row {r:2d}: {c2[:45]} -> {len(matches)} matches: {[m['sku'] for m in matches]}")
        else:
            print(f"Row {r:2d}: {c2[:45]} -> NO MATCH")
