import openpyxl
import os
import re

folder = r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES'
inv_wb = openpyxl.load_workbook(os.path.join(folder, 'INVENTARIO.xlsx'), data_only=True)
inv_ws = inv_wb.active

inv_items = []
for r in range(8, inv_ws.max_row + 1):
    sku = inv_ws.cell(r, 1).value
    name = inv_ws.cell(r, 2).value
    stock = inv_ws.cell(r, 6).value
    cost = inv_ws.cell(r, 8).value
    if sku and name:
        inv_items.append({'sku': str(sku).strip(), 'name': str(name).strip(), 'stock': stock, 'cost': cost})

wb = openpyxl.load_workbook(os.path.join(folder, 'FILTROS.xlsx'), data_only=True)

for sheet_name in ['LURICANTE MANNOL', 'LURICANTE MOBIL', 'LURICANTE AROIL', 'LURICANTE AMALIE', 'LURICANTE MOTUL', 'LURICANTE VALVOLINE']:
    ws = wb[sheet_name]
    print(f"\n=================== {sheet_name} (rows: {ws.max_row}) ===================")
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        sku_col = ws.cell(r, 1).value
        if not name and not sku_col:
            continue
        name_str = str(name).strip() if name else str(sku_col).strip()
        
        # Look for matches in inv_items
        # Extract main tokens (e.g. MANNOL TS-4 15W40 or MOBIL DELVAC MX 15W40)
        tokens = [t for t in re.split(r'[\s\-\/\(\)\,]+', name_str.upper()) if len(t) > 2]
        matches = []
        for it in inv_items:
            it_name = it['name'].upper()
            if all(t in it_name for t in tokens if t not in ['LURICANTE', 'ACEITE', 'MOTOR', 'EXTRA', 'SUPER', 'PLUS', 'SEMI', 'SINTETICO']):
                matches.append(it)
        
        print(f"Row {r:2d}: '{name_str}'")
        if matches:
            for m in matches:
                print(f"    -> SKU: {m['sku']:<16} | Stk: {str(m['stock']):>3} | {m['name']}")
        else:
            print(f"    -> NO DIRECT MATCH FOUND")
