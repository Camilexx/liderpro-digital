import openpyxl
import os

filepath = r"c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx"
wb = openpyxl.load_workbook(filepath, data_only=True)

for sheetname in wb.sheetnames:
    if sheetname == "Sheet1":
        continue
    ws = wb[sheetname]
    print(f"\n==================================================")
    print(f"SHEET: {sheetname} (max_row={ws.max_row})")
    print(f"==================================================")
    
    # Check headers (row 1 or 2)
    for r in range(1, 6):
        row_vals = [ws.cell(r, c).value for c in range(1, 8)]
        if any(v is not None and str(v).strip() != "" for v in row_vals):
            print(f"  Row {r}: {row_vals}")
