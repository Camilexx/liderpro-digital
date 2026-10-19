import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\PC\Desktop\PRODUCTOS E IMAGENES\FILTROS.xlsx', data_only=True)

for sname in wb.sheetnames:
    sheet = wb[sname]
    print(f"\n=======================================================")
    print(f"SHEET: {sname} (Rows: {sheet.max_row})")
    row_count = 0
    for r in range(1, min(sheet.max_row + 1, 40)):
        row_vals = [sheet.cell(r, c).value for c in range(1, min(sheet.max_column + 1, 6))]
        if any(row_vals):
            row_count += 1
            if row_count <= 10:
                print(f"   Row {r:2d}: {row_vals}")
    print(f"   ... Total non-empty rows inspected: {row_count}")
